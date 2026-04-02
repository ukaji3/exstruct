"""Tests for extraction pipeline fallback behavior."""

from collections.abc import Iterator
from contextlib import contextmanager
from pathlib import Path
import subprocess

from _pytest.monkeypatch import MonkeyPatch
from openpyxl import Workbook

from exstruct.core.libreoffice import LibreOfficeUnavailableError
from exstruct.core.pipeline import resolve_extraction_inputs, run_extraction_pipeline
from exstruct.errors import FallbackReason
from exstruct.models import Shape


def _make_basic_book(path: Path) -> None:
    """Create basic book for tests."""

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "v1"
    wb.save(path)
    wb.close()


def test_pipeline_fallback_skip_com_tests(
    monkeypatch: MonkeyPatch, tmp_path: Path
) -> None:
    """Verify that pipeline fallback skip COM tests."""

    path = tmp_path / "book.xlsx"
    _make_basic_book(path)
    monkeypatch.setenv("SKIP_COM_TESTS", "1")

    inputs = resolve_extraction_inputs(
        path,
        mode="standard",
        include_cell_links=False,
        include_print_areas=True,
        include_auto_page_breaks=False,
        include_colors_map=False,
        include_default_background=False,
        ignore_colors=None,
        include_formulas_map=None,
        include_merged_cells=None,
        include_merged_values_in_rows=True,
    )
    result = run_extraction_pipeline(inputs)

    assert result.state.fallback_reason == FallbackReason.SKIP_COM_TESTS
    assert result.state.com_attempted is False
    sheet = result.workbook.sheets["Sheet1"]
    assert sheet.shapes == []
    assert sheet.charts == []
    assert sheet.rows


def test_pipeline_fallback_com_unavailable(
    monkeypatch: MonkeyPatch, tmp_path: Path
) -> None:
    """
    Verifies that the extraction pipeline falls back when COM access is unavailable.

    Creates a basic workbook, forces the COM-access entry point to raise, runs the extraction pipeline, and asserts that the pipeline records a fallback due to COM being unavailable (`FallbackReason.COM_UNAVAILABLE`), did not attempt COM (`com_attempted is False`), and that the resulting sheet "Sheet1" exists, contains rows, and has no shapes or charts.
    """
    path = tmp_path / "book.xlsx"
    _make_basic_book(path)
    monkeypatch.delenv("SKIP_COM_TESTS", raising=False)

    def _raise(*_args: object, **_kwargs: object) -> None:
        """Raise the expected test exception."""

        raise RuntimeError("no COM")

    monkeypatch.setattr("exstruct.core.pipeline.xlwings_workbook", _raise)

    inputs = resolve_extraction_inputs(
        path,
        mode="standard",
        include_cell_links=False,
        include_print_areas=True,
        include_auto_page_breaks=False,
        include_colors_map=False,
        include_default_background=False,
        ignore_colors=None,
        include_formulas_map=None,
        include_merged_cells=None,
        include_merged_values_in_rows=True,
    )
    result = run_extraction_pipeline(inputs)

    assert result.state.fallback_reason == FallbackReason.COM_UNAVAILABLE
    assert result.state.com_attempted is False
    sheet = result.workbook.sheets["Sheet1"]
    assert sheet.shapes == []
    assert sheet.charts == []
    assert sheet.rows


def test_pipeline_fallback_com_pipeline_failed(
    monkeypatch: MonkeyPatch, tmp_path: Path
) -> None:
    """Verify that pipeline fallback COM pipeline failed."""

    path = tmp_path / "book.xlsx"
    _make_basic_book(path)
    monkeypatch.delenv("SKIP_COM_TESTS", raising=False)

    @contextmanager
    def _dummy_workbook(_path: Path) -> Iterator[object]:
        """Yield a dummy workbook context manager for this test."""

        yield object()

    def _raise(
        *_args: object,
        **_kwargs: object,
    ) -> None:
        """Raise the expected test exception."""

        raise RuntimeError("pipeline failed")

    monkeypatch.setattr("exstruct.core.pipeline.xlwings_workbook", _dummy_workbook)
    monkeypatch.setattr("exstruct.core.pipeline.run_com_pipeline", _raise)

    inputs = resolve_extraction_inputs(
        path,
        mode="standard",
        include_cell_links=False,
        include_print_areas=True,
        include_auto_page_breaks=False,
        include_colors_map=False,
        include_default_background=False,
        ignore_colors=None,
        include_formulas_map=None,
        include_merged_cells=None,
        include_merged_values_in_rows=True,
    )
    result = run_extraction_pipeline(inputs)

    assert result.state.fallback_reason == FallbackReason.COM_PIPELINE_FAILED
    assert result.state.com_attempted is True
    sheet = result.workbook.sheets["Sheet1"]
    assert sheet.shapes == []
    assert sheet.charts == []
    assert sheet.rows


def test_pipeline_fallback_libreoffice_unavailable(
    monkeypatch: MonkeyPatch, tmp_path: Path
) -> None:
    """Verify that pipeline fallback LibreOffice unavailable."""

    path = tmp_path / "book.xlsx"
    _make_basic_book(path)

    def _raise(**_kwargs: object) -> object:
        """Raise the expected test exception."""

        raise LibreOfficeUnavailableError("missing soffice")

    monkeypatch.setattr("exstruct.core.pipeline.resolve_rich_backend", _raise)

    inputs = resolve_extraction_inputs(
        path,
        mode="libreoffice",
        include_cell_links=False,
        include_print_areas=True,
        include_auto_page_breaks=False,
        include_colors_map=False,
        include_default_background=False,
        ignore_colors=None,
        include_formulas_map=None,
        include_merged_cells=None,
        include_merged_values_in_rows=True,
    )
    result = run_extraction_pipeline(inputs)

    assert result.state.fallback_reason == FallbackReason.LIBREOFFICE_UNAVAILABLE
    sheet = result.workbook.sheets["Sheet1"]
    assert sheet.shapes == []
    assert sheet.charts == []
    assert sheet.rows


def test_pipeline_fallback_libreoffice_pipeline_failed(
    monkeypatch: MonkeyPatch, tmp_path: Path
) -> None:
    """Verify that pipeline fallback LibreOffice pipeline failed."""

    path = tmp_path / "book.xlsx"
    _make_basic_book(path)

    class _Backend:
        """Backend test double used by pipeline and mode tests."""

        def extract_shapes(self, *, mode: str) -> dict[str, list[object]]:
            """Provide the shape-extraction behavior for this test double."""

            _ = mode
            raise RuntimeError("boom")

        def extract_charts(self, *, mode: str) -> dict[str, list[object]]:
            """Provide the chart-extraction behavior for this test double."""

            _ = mode
            return {}

    monkeypatch.setattr(
        "exstruct.core.pipeline.resolve_rich_backend",
        lambda **_kwargs: _Backend(),
    )

    inputs = resolve_extraction_inputs(
        path,
        mode="libreoffice",
        include_cell_links=False,
        include_print_areas=True,
        include_auto_page_breaks=False,
        include_colors_map=False,
        include_default_background=False,
        ignore_colors=None,
        include_formulas_map=None,
        include_merged_cells=None,
        include_merged_values_in_rows=True,
    )
    result = run_extraction_pipeline(inputs)

    assert result.state.fallback_reason == FallbackReason.LIBREOFFICE_PIPELINE_FAILED
    sheet = result.workbook.sheets["Sheet1"]
    assert sheet.shapes == []
    assert sheet.charts == []
    assert sheet.rows


def test_pipeline_fallback_libreoffice_preserves_shapes_when_chart_extraction_fails(
    monkeypatch: MonkeyPatch, tmp_path: Path
) -> None:
    """Verify that chart failures keep already extracted shapes in the workbook."""

    path = tmp_path / "book.xlsx"
    _make_basic_book(path)

    class _Backend:
        """Backend test double used by pipeline fallback tests."""

        def extract_shapes(self, *, mode: str) -> dict[str, list[Shape]]:
            _ = mode
            return {"Sheet1": [Shape(id=1, text="shape", l=0, t=0)]}

        def extract_charts(self, *, mode: str) -> dict[str, list[object]]:
            _ = mode
            raise RuntimeError("chart boom")

    monkeypatch.setattr(
        "exstruct.core.pipeline.resolve_rich_backend",
        lambda **_kwargs: _Backend(),
    )

    inputs = resolve_extraction_inputs(
        path,
        mode="libreoffice",
        include_cell_links=False,
        include_print_areas=True,
        include_auto_page_breaks=False,
        include_colors_map=False,
        include_default_background=False,
        ignore_colors=None,
        include_formulas_map=None,
        include_merged_cells=None,
        include_merged_values_in_rows=True,
    )
    result = run_extraction_pipeline(inputs)

    assert result.state.fallback_reason == FallbackReason.LIBREOFFICE_PIPELINE_FAILED
    sheet = result.workbook.sheets["Sheet1"]
    assert len(sheet.shapes) == 1
    assert sheet.shapes[0].text == "shape"
    assert sheet.charts == []
    assert sheet.rows


def test_pipeline_fallback_libreoffice_shape_failure_short_circuits_charts(
    monkeypatch: MonkeyPatch, tmp_path: Path
) -> None:
    """Verify that shape failures do not continue into chart extraction."""

    path = tmp_path / "book.xlsx"
    _make_basic_book(path)
    chart_calls: list[str] = []

    class _Backend:
        """Backend test double used by pipeline fallback tests."""

        def extract_shapes(self, *, mode: str) -> dict[str, list[Shape]]:
            _ = mode
            raise RuntimeError("shape boom")

        def extract_charts(self, *, mode: str) -> dict[str, list[object]]:
            _ = mode
            chart_calls.append("called")
            return {}

    monkeypatch.setattr(
        "exstruct.core.pipeline.resolve_rich_backend",
        lambda **_kwargs: _Backend(),
    )

    inputs = resolve_extraction_inputs(
        path,
        mode="libreoffice",
        include_cell_links=False,
        include_print_areas=True,
        include_auto_page_breaks=False,
        include_colors_map=False,
        include_default_background=False,
        ignore_colors=None,
        include_formulas_map=None,
        include_merged_cells=None,
        include_merged_values_in_rows=True,
    )
    result = run_extraction_pipeline(inputs)

    assert result.state.fallback_reason == FallbackReason.LIBREOFFICE_PIPELINE_FAILED
    assert chart_calls == []
    sheet = result.workbook.sheets["Sheet1"]
    assert sheet.shapes == []
    assert sheet.charts == []
    assert sheet.rows


def test_pipeline_fallback_libreoffice_incompatible_override(
    monkeypatch: MonkeyPatch, tmp_path: Path
) -> None:
    """Verify that an incompatible override is treated as LibreOffice unavailable."""

    path = tmp_path / "book.xlsx"
    _make_basic_book(path)
    soffice_path = tmp_path / "soffice"
    soffice_path.write_text("", encoding="utf-8")
    override_path = tmp_path / "python3"
    override_path.write_text("", encoding="utf-8")
    monkeypatch.setenv("EXSTRUCT_LIBREOFFICE_PATH", str(soffice_path))
    monkeypatch.setenv("EXSTRUCT_LIBREOFFICE_PYTHON_PATH", str(override_path))

    def _fake_run(
        *_args: object, **_kwargs: object
    ) -> subprocess.CompletedProcess[str]:
        """Simulate bridge probe failure for the configured override."""

        raise subprocess.CalledProcessError(
            returncode=1,
            cmd=str(override_path),
            stderr="SyntaxError: invalid syntax",
        )

    monkeypatch.setattr("exstruct.core.libreoffice.subprocess.run", _fake_run)

    inputs = resolve_extraction_inputs(
        path,
        mode="libreoffice",
        include_cell_links=False,
        include_print_areas=True,
        include_auto_page_breaks=False,
        include_colors_map=False,
        include_default_background=False,
        ignore_colors=None,
        include_formulas_map=None,
        include_merged_cells=None,
        include_merged_values_in_rows=True,
    )
    result = run_extraction_pipeline(inputs)

    assert result.state.fallback_reason == FallbackReason.LIBREOFFICE_UNAVAILABLE
    sheet = result.workbook.sheets["Sheet1"]
    assert sheet.shapes == []
    assert sheet.charts == []
    assert sheet.rows
